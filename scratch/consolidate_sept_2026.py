"""
Consolidate SEPT 2026 Conversion Data from regional sheets (TN, AP_TS, KL)
into standardized Excel templates with data validations, auto-calculating formulas,
and executive KPI summaries.
"""

import os
import sys
import re

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
import numpy as np

# =============================================================================
# STANDARDIZATION DICTIONARIES
# =============================================================================

MAKE_MAP = {
    'maruti': 'MARUTI SUZUKI', 'maruti suzuki': 'MARUTI SUZUKI',
    'maruthi': 'MARUTI SUZUKI', 'maruthi suzuki': 'MARUTI SUZUKI',
    'hyundai': 'HYUNDAI',
    'tata': 'TATA',
    'toyota': 'TOYOTA', 'totoya': 'TOYOTA',
    'honda': 'HONDA',
    'mahindra': 'MAHINDRA',
    'renault': 'RENAULT', 'renalut': 'RENAULT',
    'volkswagen': 'VOLKSWAGEN', 'vw': 'VOLKSWAGEN',
    'ford': 'FORD', 'ford figo': 'FORD',
    'nissan': 'NISSAN',
    'datsun': 'DATSUN', 'datson': 'DATSUN',
    'skoda': 'SKODA',
    'kia': 'KIA',
    'mg': 'MG MOTORS', 'mg motors': 'MG MOTORS',
    'chevrolet': 'CHEVROLET', 'cheverolet': 'CHEVROLET',
    'fiat': 'FIAT',
    'citroen': 'CITROEN', 'cetrion': 'CITROEN',
    'bharatbenz': 'BHARATBENZ',
    'hindustan motors': 'HINDUSTAN MOTORS',
    'jeep': 'JEEP'
}

MODEL_MAP = {
    'alt0 800': 'Alto 800', 'alto 800': 'Alto 800', 'alto800': 'Alto 800',
    'alto k10': 'Alto K10', 'altok10': 'Alto K10', 'k 10': 'Alto K10', 'k10': 'Alto K10',
    'alto': 'Alto', '800': '800', 'm 800': '800', 'm800': '800',
    'a-star': 'A-Star', 'astar': 'A-Star', 'a star': 'A-Star',
    'wagon r': 'Wagon R', 'wagonr': 'Wagon R', 'wagon-r': 'Wagon R',
    'swift': 'Swift',
    'dzire': 'Dzire', 'dezire': 'Dzire', 'swift dzire': 'Dzire', 'swiftdzire': 'Dzire', 'swift desire': 'Dzire',
    'baleno': 'Baleno', 'celerio': 'Celerio', 'ciaz': 'Ciaz',
    'breeza': 'Vitara Brezza', 'brezza': 'Vitara Brezza', 'vitara brezza': 'Vitara Brezza',
    'grand vitara': 'Grand Vitara', 'eeco': 'Eeco', 'ertiga': 'Ertiga', 'esteem': 'Esteem',
    'estilo': 'Zen Estilo', 'zen estilo': 'Zen Estilo', 'fronx': 'Fronx', 'ignis': 'Ignis',
    'omni': 'Omni', 'ritz': 'Ritz', 'scross': 'S-Cross', 'sx4': 'SX4', 'xl 6': 'XL6', 'zen': 'Zen',
    'grand i10': 'Grand i10', 'grand': 'Grand i10', 'g i10': 'Grand i10', 'grand i10 nios': 'Grand i10 Nios',
    'i10': 'i10', 'i 10': 'i10', 'i20': 'i20', 'i 20': 'i20', 'i20 active': 'i20 Active',
    'eon': 'Eon', 'santro': 'Santro', 'santafe': 'Santa Fe', 'venue': 'Venue', 'verna': 'Verna',
    'xcent': 'Xcent', 'creta': 'Creta', 'tucson': 'Tucson', 'accent': 'Accent',
    'tiago': 'Tiago', 'tigor': 'Tigor', 'tigore': 'Tigor', 'altroz': 'Altroz', 'bolt': 'Bolt',
    'indica': 'Indica', 'indigo': 'Indigo', 'indigo cs': 'Indigo', 'nano': 'Nano', 'nexon': 'Nexon',
    'punch': 'Punch', 'safari': 'Safari', 'zest': 'Zest', 'ultra': 'Ultra',
    'city': 'City', 'civic': 'Civic', 'amaze': 'Amaze', 'accord': 'Accord', 'brio': 'Brio', 'jazz': 'Jazz', 'wr-v': 'WR-V',
    'etios': 'Etios', 'etios liva': 'Etios Liva', 'glanza': 'Glanza', 'hyryder': 'Hyryder', 'corolla': 'Corolla Altis',
    'figo': 'Figo', 'ecosport': 'EcoSport', 'fiesta': 'Fiesta', 'fiestaclassic': 'Fiesta Classic',
    'bolero': 'Bolero', 'scorpio': 'Scorpio', 'marazzo': 'Marazzo', 'kuv 100': 'KUV100', 'xuv300': 'XUV300',
    'xuv 500': 'XUV500', 'xuv 700': 'XUV700', 'verito': 'Verito',
    'kwid': 'Kwid', 'duster': 'Duster', 'triber': 'Triber', 'kiger': 'Kiger', 'pulse': 'Pulse', 'scala': 'Scala',
    'polo': 'Polo', 'vento': 'Vento', 'ameo': 'Ameo', 'rapid': 'Rapid', 'kushaq': 'Kushaq',
    'octavia': 'Octavia', 'laura': 'Laura', 'slavia': 'Slavia', 'fabia': 'Fabia',
    'micra': 'Micra', 'sunny': 'Sunny', 'magnite': 'Magnite', 'terrano': 'Terrano',
    'datsun go': 'Go', 'go': 'Go', 'redi-go': 'Redi-GO', 'redi-g0': 'Redi-GO', 'redigo': 'Redi-GO',
    'seltos': 'Seltos', 'sonet': 'Sonet',
    'beat': 'Beat', 'spark': 'Spark', 'cruze': 'Cruze', 'aveo': 'Aveo',
    'punto': 'Punto', 'grand punto': 'Grand Punto', 'palio stile': 'Palio Stile',
    'compass': 'Compass', 'zs ev': 'ZS EV', 'c3': 'C3', 'ambassador': 'Ambassador', '1017': '1017'
}

SELLER_MAP = {
    'castel traders honada': 'CASTLE TRADERS HONDA',
    'kalyanimotors maruthi': 'KALYANI MOTORS MARUTI',
    'popular motor world': 'POPULAR MOTORWORLD PVT LTD',
    'popular vehicles and services limited': 'POPULAR VEHICLES',
    'krish toyota - coimbatore': 'KRISH TOYOTA',
    'castle traders maruti': 'CASTLE TRADERS MARUTI',
    'daimler inida commercial': 'DAIMLER INDIA COMMERCIAL',
    'em cetrion': 'EVM CITROEN',
    'evm citron': 'EVM CITROEN',
    'citroen evm': 'EVM CITROEN',
    'am moters': 'AM MOTORS',
    'apco  hyundai': 'APCO HYUNDAI',
    'kvr  tata': 'KVR TATA',
    'marvel motots': 'MARVEL MOTORS',
    'incheon kia': 'INCHEON KIA',
    'pinnacle nissan': 'PINNACLE NISSAN',
    'luxon tata': 'LUXON TATA',
    'gokulam tata': 'GOKULAM TATA',
    'kairali toyota': 'KAIRALI TOYOTA',
}

BUYER_MAP = {
    'cars 24': 'Cars24', 'cars24': 'Cars24',
    'auto bot': 'Autobot', 'autoboot': 'Autobot', 'autobot': 'Autobot',
    'aito cars': 'Aitos Cars', 'aitos cars': 'Aitos Cars',
    'ulavan enterprises': 'Uzhavan Enterprise', 'uzhavan enterprise': 'Uzhavan Enterprise',
    'deream drive': 'Dream Drive', 'dream drive': 'Dream Drive',
    '33cars': '33 Cars',
    'ksm auto consulting': 'KSM Auto Consulting', 'kam auto conselting': 'KSM Auto Consulting',
    'sarvesh caars': 'Sarvesh Cars', 'sarvesh cars': 'Sarvesh Cars',
    'shri shakthi vinayaga cars': 'Sri Shakthi Vinayaga Cars',
    'm/s yes cars': 'Yes Cars', 'yes cars': 'Yes Cars',
    'vs enterprises': 'VS Enterprises',
    'thiya cars': 'Thiya Cars',
    'santhosh cars': 'Santhosh Cars',
    'nirmal cars': 'Nirmal Cars',
    'grid 33': 'Grid 33',
    'popular thrissur': 'Popular Thrissur',
}

def clean_vehicle_number(val):
    if not val or pd.isna(val):
        return ""
    return str(val).strip().replace(" ", "").upper()

def standardize_make(val, model_val=None):
    if not val or pd.isna(val):
        return ""
    s = str(val).strip().lower()
    if s == 'ford figo':
        return 'FORD'
    return MAKE_MAP.get(s, str(val).strip().upper())

def standardize_model(val, make_val=None):
    if not val or pd.isna(val):
        return ""
    s = str(val).strip().lower()
    if s in MODEL_MAP:
        return MODEL_MAP[s]
    # Clean non-alphanumeric
    k_clean = re.sub(r'[^a-z0-9]', '', s)
    for k, tgt in MODEL_MAP.items():
        if re.sub(r'[^a-z0-9]', '', k) == k_clean:
            return tgt
    return str(val).strip().title()

def standardize_seller(val):
    if not val or pd.isna(val):
        return ""
    s = str(val).strip().lower()
    return SELLER_MAP.get(s, str(val).strip().upper())

def derive_seller_segment(seller_name, explicit_seg=None):
    if explicit_seg and not pd.isna(explicit_seg):
        s_seg = str(explicit_seg).strip().lower()
        if 'non' in s_seg:
            return "Non-Captive"
        if 'captive' in s_seg:
            return "Captive"
    if not seller_name or pd.isna(seller_name):
        return "Non-Captive"
    s = str(seller_name).strip().lower()
    for kw in ["tvs mobility", "exchange", "tvs ", "certified internal"]:
        if kw in s:
            return "Captive"
    return "Non-Captive"

def standardize_buyer(val):
    if not val or pd.isna(val):
        return ""
    s = str(val).strip()
    k = s.lower()
    if k in BUYER_MAP:
        return BUYER_MAP[k]
    # Format nicely with Title Case
    return s.title()

def parse_num(val, is_int=False):
    if val is None or pd.isna(val):
        return None
    try:
        f = float(str(val).replace(",", "").replace(" ", "").strip())
        return int(round(f)) if is_int else round(f, 2)
    except (ValueError, TypeError):
        return None

def safe_save(wb, output_path):
    try:
        wb.save(output_path)
        print(f"  -> Successfully saved: {output_path}")
    except PermissionError:
        alt = os.path.splitext(output_path)[0] + "_v2.xlsx"
        try:
            wb.save(alt)
            print(f"  -> File is locked by Excel; saved copy to: {alt}")
        except Exception as e:
            print(f"  -> Could not save {output_path}: {e}")


# =============================================================================
# EXTRACT DATA FROM SEPT 2026.xlsx
# =============================================================================

def load_sept_data(filepath):
    print(f"Loading raw September data from: {filepath}")
    records = []

    # 1. TN Sheet
    df_tn = pd.read_excel(filepath, sheet_name='TN')
    for _, row in df_tn.iterrows():
        veh_no = clean_vehicle_number(row['REG Number'])
        if not veh_no:
            continue
        
        mfg_year = parse_num(row.get('MFG Year'), is_int=True)
        if mfg_year is None:
            # check reg year
            reg_val = row.get('REG Year')
            if pd.notna(reg_val):
                if isinstance(reg_val, (datetime, pd.Timestamp)):
                    mfg_year = reg_val.year
                else:
                    try:
                        mfg_year = int(str(reg_val)[:4])
                    except:
                        pass
        
        odo = parse_num(row.get('KM Driven'), is_int=True)
        bid = parse_num(row.get('BID'), is_int=False)
        inc_gst = parse_num(row.get('Total Margin'), is_int=False)
        ex_gst = parse_num(row.get('Net Margin'), is_int=False)
        
        seller = standardize_seller(row.get('Seller Name'))
        seg = derive_seller_segment(seller)
        buyer = standardize_buyer(row.get('Sold To'))
        
        mk_raw = row.get('Make')
        md_raw = row.get('Model')
        make = standardize_make(mk_raw)
        model = standardize_model(md_raw, make)
        
        records.append({
            'FY': 'FY27',
            'Month': 'Sep',
            'State': 'TN',
            'Vehicle Number': veh_no,
            'Make': make,
            'Model': model,
            'Seller Name': seller,
            'Seller Segment': seg,
            'Buyer Name': buyer,
            'Year': mfg_year,
            'Odometer': odo,
            'Final Bid Value': bid,
            'Total Income (Ex GST)': ex_gst,
            'Total Income (Inc GST)': inc_gst,
            'GST (18%)': round(inc_gst - ex_gst, 2) if (inc_gst and ex_gst) else None
        })
    print(f"  -> Extracted TN: {len(df_tn)} records")

    # 2. AP_TS Sheet
    df_ap = pd.read_excel(filepath, sheet_name='AP_TS', header=None).iloc[2:]
    for _, row in df_ap.iterrows():
        veh_no = clean_vehicle_number(row[4])
        if not veh_no:
            continue
        
        ex_gst = parse_num(row[13], is_int=False)
        inc_gst = round(ex_gst * 1.18, 2) if ex_gst is not None else None
        gst_amt = round(inc_gst - ex_gst, 2) if (inc_gst and ex_gst) else None
        
        seller_raw = row[10]
        seller = str(seller_raw).strip().upper() if pd.notna(seller_raw) else ""
        seg = derive_seller_segment(seller, row[11])
        buyer = standardize_buyer(row[12])
        
        make = standardize_make(row[5])
        model = standardize_model(row[6], make)
        
        records.append({
            'FY': 'FY27',
            'Month': 'Sep',
            'State': 'AP/TS',
            'Vehicle Number': veh_no,
            'Make': make,
            'Model': model,
            'Seller Name': seller,
            'Seller Segment': seg,
            'Buyer Name': buyer,
            'Year': parse_num(row[7], is_int=True),
            'Odometer': parse_num(row[8], is_int=True),
            'Final Bid Value': parse_num(row[9], is_int=False),
            'Total Income (Ex GST)': ex_gst,
            'Total Income (Inc GST)': inc_gst,
            'GST (18%)': gst_amt
        })
    print(f"  -> Extracted AP/TS: {len(df_ap)} records")

    # 3. KL Sheet
    df_kl = pd.read_excel(filepath, sheet_name='KL', header=None).iloc[3:]
    for _, row in df_kl.iterrows():
        veh_no = clean_vehicle_number(row[5])
        if not veh_no:
            continue
        
        inc_gst = parse_num(row[14], is_int=False)
        ex_gst = parse_num(row[15], is_int=False)
        gst_amt = round(inc_gst - ex_gst, 2) if (inc_gst and ex_gst) else None
        
        seller = standardize_seller(row[11])
        seg = derive_seller_segment(seller, row[12])
        buyer = standardize_buyer(row[13])
        
        make = standardize_make(row[6])
        model = standardize_model(row[7], make)
        
        records.append({
            'FY': 'FY27',
            'Month': 'Sep',
            'State': 'Kerala',
            'Vehicle Number': veh_no,
            'Make': make,
            'Model': model,
            'Seller Name': seller,
            'Seller Segment': seg,
            'Buyer Name': buyer,
            'Year': parse_num(row[8], is_int=True),
            'Odometer': parse_num(row[9], is_int=True),
            'Final Bid Value': parse_num(row[10], is_int=False),
            'Total Income (Ex GST)': ex_gst,
            'Total Income (Inc GST)': inc_gst,
            'GST (18%)': gst_amt
        })
    print(f"  -> Extracted Kerala: {len(df_kl)} records")

    # Sort canonically: AP/TS, Kerala, TN
    state_order = {"AP/TS": 1, "Kerala": 2, "TN": 3}
    records.sort(key=lambda r: (state_order.get(r['State'], 99), r['Vehicle Number']))
    print(f"Total Consolidated September Records: {len(records)}")
    return records


# =============================================================================
# WRITER 1: STANDARD 13-COLUMN MANAGER TEMPLATE
# (Matches Consolidated_Conversion_Data_Aug26_Final.xlsx)
# =============================================================================

def write_standard_13col_template(records, output_path):
    print(f"Writing Standard 13-Column Template: {os.path.basename(output_path)}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conversion Data"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"),
    )
    data_font = Font(name="Calibri", size=10)
    number_format_comma = '#,##0'
    number_format_curr = '#,##0.00'

    # Title Row
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = "Only Conversion car"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Subtitle Row
    ws.merge_cells("A2:M2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = "ex. GST"
    subtitle_cell.font = Font(name="Calibri", size=10, italic=True, color="808080")
    subtitle_cell.alignment = Alignment(horizontal="right", vertical="center")

    columns = [
        "FY", "Month", "State", "Vehicle Number", "Make", "Model",
        "Year", "Odometer", "Final Bid Value", "Seller Name",
        "Seller Segment", "Buyer Name", "Total Income"
    ]

    for col_idx, h in enumerate(columns, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    light_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

    for r_idx, rec in enumerate(records, start=4):
        vals = [
            rec['FY'], rec['Month'], rec['State'], rec['Vehicle Number'],
            rec['Make'], rec['Model'], rec['Year'], rec['Odometer'],
            rec['Final Bid Value'], rec['Seller Name'], rec['Seller Segment'],
            rec['Buyer Name'], rec['Total Income (Ex GST)']
        ]
        is_even = (r_idx % 2 == 0)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = light_fill
            
            if c_idx in (1, 2, 3):
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 7: # Year
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "0"
            elif c_idx in (8, 9): # Odo, Bid
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = number_format_comma
            elif c_idx == 13: # Total Income (Ex GST)
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = number_format_curr

    # Column dimensions
    col_widths = {1: 10, 2: 10, 3: 12, 4: 18, 5: 20, 6: 20, 7: 10,
                  8: 15, 9: 18, 10: 36, 11: 16, 12: 34, 13: 20}
    for c_i, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_i)].width = w

    # Make list sheet for validation
    make_sheet = wb.create_sheet("_MakeList")
    make_sheet.sheet_state = "hidden"
    unique_makes = sorted(list(set(MAKE_MAP.values())))
    for i, m in enumerate(unique_makes, start=1):
        make_sheet.cell(row=i, column=1, value=m)

    max_row = max(len(records) + 50, 1000)

    # Validations
    dv_fy = DataValidation(type="list", formula1='"FY25,FY26,FY27,FY28"', allow_blank=True)
    dv_fy.add(f"A4:A{max_row}")
    ws.add_data_validation(dv_fy)

    dv_month = DataValidation(type="list", formula1='"Jan,Feb,Mar,Apr,May,Jun,Jul,Aug,Sep,Oct,Nov,Dec"', allow_blank=True)
    dv_month.add(f"B4:B{max_row}")
    ws.add_data_validation(dv_month)

    dv_state = DataValidation(type="list", formula1='"AP/TS,Kerala,TN"', allow_blank=True)
    dv_state.add(f"C4:C{max_row}")
    ws.add_data_validation(dv_state)

    dv_segment = DataValidation(type="list", formula1='"Captive,Non-Captive"', allow_blank=True)
    dv_segment.add(f"K4:K{max_row}")
    ws.add_data_validation(dv_segment)

    dv_make = DataValidation(type="list", formula1=f"=_MakeList!$A$1:$A${len(unique_makes)}", allow_blank=True)
    dv_make.add(f"E4:E{max_row}")
    ws.add_data_validation(dv_make)

    safe_save(wb, output_path)


# =============================================================================
# WRITER 2: FULL EXECUTIVE ANALYTICS WORKBOOK
# (Matches Consolidated_Conversion_Data_Aug26_Analytics.xlsx)
# =============================================================================

def write_analytics_workbook(records, output_path, subtitle="September 2026"):
    print(f"Writing Executive Analytics Workbook: {os.path.basename(output_path)}")
    wb = openpyxl.Workbook()

    # Sheet 1: Executive Analytics & KPIs
    ws_kpi = wb.active
    ws_kpi.title = "Executive Analytics & KPIs"
    ws_kpi.views.sheetView[0].showGridLines = True

    # Branding header
    ws_kpi.merge_cells("B1:J1")
    ws_kpi["B1"] = "TVS Certified - Conversion Analytics & Performance Intelligence"
    ws_kpi["B1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws_kpi["B1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_kpi.merge_cells("B2:J2")
    ws_kpi["B2"] = f"{subtitle} Executive Summary & Regional Conversion Metrics ({len(records):,} Units)"
    ws_kpi["B2"].font = Font(name="Calibri", size=11, bold=True, color="475569")
    ws_kpi["B2"].alignment = Alignment(horizontal="center", vertical="center")

    kpi_cards = [
        ("B4", "TOTAL CONVERTED", f"=COUNTA('Conversion Analytics (Master)'!D4:D{len(records)+3})", "#,##0"),
        ("D4", "AUCTION TURNOVER", f"=SUM('Conversion Analytics (Master)'!I4:I{len(records)+3})", "₹#,##0.00"),
        ("F4", "TVS INCOME (EX-GST)", f"=SUM('Conversion Analytics (Master)'!M4:M{len(records)+3})", "₹#,##0.00"),
        ("H4", "TVS INCOME (INCL-GST)", f"=SUM('Conversion Analytics (Master)'!N4:N{len(records)+3})", "₹#,##0.00"),
        ("J4", "AVG TVS MARGIN %", f"=AVERAGE('Conversion Analytics (Master)'!O4:O{len(records)+3})", "0.00%"),
    ]

    card_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    card_value_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    border_box = Border(
        left=Side(style="thin", color="94A3B8"), right=Side(style="thin", color="94A3B8"),
        top=Side(style="thin", color="94A3B8"), bottom=Side(style="thin", color="94A3B8"),
    )

    for col_pos, title, form, fmt in kpi_cards:
        col_letter = col_pos[0]
        # Label cell
        cell_lbl = ws_kpi[f"{col_letter}4"]
        cell_lbl.value = title
        cell_lbl.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell_lbl.fill = card_header_fill
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.border = border_box

        # Value cell
        cell_val = ws_kpi[f"{col_letter}5"]
        cell_val.value = form
        cell_val.font = Font(name="Calibri", size=12, bold=True, color="0F172A")
        cell_val.fill = card_value_fill
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.border = border_box
        cell_val.number_format = fmt

    # Regional Performance Table
    ws_kpi["B8"] = "REGIONAL PERFORMANCE BREAKDOWN"
    ws_kpi["B8"].font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")

    reg_headers = ["State", "Units Converted", "GMV (Turnover)", "TVS Income (Ex GST)", "TVS Income (Inc GST)", "Avg Bid Value", "Avg Take Rate %"]
    for c_i, h in enumerate(reg_headers, start=2):
        col_let = get_column_letter(c_i)
        c = ws_kpi[f"{col_let}9"]
        c.value = h
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_box

    states = ["AP/TS", "Kerala", "TN"]
    last_r = len(records) + 3
    for s_idx, st in enumerate(states, start=10):
        # State
        ws_kpi[f"B{s_idx}"] = st
        ws_kpi[f"B{s_idx}"].alignment = Alignment(horizontal="center")
        ws_kpi[f"B{s_idx}"].border = border_box

        # Units
        ws_kpi[f"C{s_idx}"] = f"=COUNTIF('Conversion Analytics (Master)'!$C$4:$C${last_r}, B{s_idx})"
        ws_kpi[f"C{s_idx}"].number_format = "#,##0"
        ws_kpi[f"C{s_idx}"].border = border_box

        # GMV
        ws_kpi[f"D{s_idx}"] = f"=SUMIF('Conversion Analytics (Master)'!$C$4:$C${last_r}, B{s_idx}, 'Conversion Analytics (Master)'!$I$4:$I${last_r})"
        ws_kpi[f"D{s_idx}"].number_format = "₹#,##0.00"
        ws_kpi[f"D{s_idx}"].border = border_box

        # Income Ex GST
        ws_kpi[f"E{s_idx}"] = f"=SUMIF('Conversion Analytics (Master)'!$C$4:$C${last_r}, B{s_idx}, 'Conversion Analytics (Master)'!$M$4:$M${last_r})"
        ws_kpi[f"E{s_idx}"].number_format = "₹#,##0.00"
        ws_kpi[f"E{s_idx}"].border = border_box

        # Income Inc GST
        ws_kpi[f"F{s_idx}"] = f"=SUMIF('Conversion Analytics (Master)'!$C$4:$C${last_r}, B{s_idx}, 'Conversion Analytics (Master)'!$N$4:$N${last_r})"
        ws_kpi[f"F{s_idx}"].number_format = "₹#,##0.00"
        ws_kpi[f"F{s_idx}"].border = border_box

        # Avg Bid
        ws_kpi[f"G{s_idx}"] = f'=IF(C{s_idx}>0, D{s_idx}/C{s_idx}, 0)'
        ws_kpi[f"G{s_idx}"].number_format = "₹#,##0.00"
        ws_kpi[f"G{s_idx}"].border = border_box

        # Take Rate %
        ws_kpi[f"H{s_idx}"] = f'=IF(D{s_idx}>0, E{s_idx}/D{s_idx}, 0)'
        ws_kpi[f"H{s_idx}"].number_format = "0.00%"
        ws_kpi[f"H{s_idx}"].border = border_box

    # Total Row
    tot_row = 13
    ws_kpi[f"B{tot_row}"] = "TOTAL"
    ws_kpi[f"B{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"B{tot_row}"].border = border_box
    ws_kpi[f"B{tot_row}"].alignment = Alignment(horizontal="center")

    ws_kpi[f"C{tot_row}"] = f"=SUM(C10:C12)"
    ws_kpi[f"C{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"C{tot_row}"].number_format = "#,##0"
    ws_kpi[f"C{tot_row}"].border = border_box

    ws_kpi[f"D{tot_row}"] = f"=SUM(D10:D12)"
    ws_kpi[f"D{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"D{tot_row}"].number_format = "₹#,##0.00"
    ws_kpi[f"D{tot_row}"].border = border_box

    ws_kpi[f"E{tot_row}"] = f"=SUM(E10:E12)"
    ws_kpi[f"E{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"E{tot_row}"].number_format = "₹#,##0.00"
    ws_kpi[f"E{tot_row}"].border = border_box

    ws_kpi[f"F{tot_row}"] = f"=SUM(F10:F12)"
    ws_kpi[f"F{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"F{tot_row}"].number_format = "₹#,##0.00"
    ws_kpi[f"F{tot_row}"].border = border_box

    ws_kpi[f"G{tot_row}"] = f'=IF(C{tot_row}>0, D{tot_row}/C{tot_row}, 0)'
    ws_kpi[f"G{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"G{tot_row}"].number_format = "₹#,##0.00"
    ws_kpi[f"G{tot_row}"].border = border_box

    ws_kpi[f"H{tot_row}"] = f'=IF(D{tot_row}>0, E{tot_row}/D{tot_row}, 0)'
    ws_kpi[f"H{tot_row}"].font = Font(name="Calibri", size=10, bold=True)
    ws_kpi[f"H{tot_row}"].number_format = "0.00%"
    ws_kpi[f"H{tot_row}"].border = border_box

    for col in ["B", "C", "D", "E", "F", "G", "H", "I", "J"]:
        ws_kpi.column_dimensions[col].width = 22

    # Sheet 2: Conversion Analytics (Master)
    ws_master = wb.create_sheet("Conversion Analytics (Master)")
    ws_master.views.sheetView[0].showGridLines = True
    ws_master.freeze_panes = "A4"

    ws_master.merge_cells("A1:O1")
    ws_master["A1"] = "TVS Certified - Conversion Analytics Master Dataset"
    ws_master["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws_master["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_master.merge_cells("A2:O2")
    ws_master["A2"] = f"{subtitle} Analytical Inventory Profile ({len(records):,} Converted Units)"
    ws_master["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")
    ws_master["A2"].alignment = Alignment(horizontal="center", vertical="center")

    m_headers = [
        "FY", "Month", "State", "Vehicle Number", "Make", "Model",
        "Year", "Odometer", "Final Bid Value", "Seller Name",
        "Seller Segment", "Buyer Name", "Total Income (ex. GST)", "Total Income", "TVS Margin %"
    ]

    header_font_m = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill_m = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    border_m = Border(
        left=Side(style="thin", color="CBD5E1"), right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"), bottom=Side(style="thin", color="CBD5E1")
    )

    for c_i, h in enumerate(m_headers, start=1):
        c = ws_master.cell(row=3, column=c_i, value=h)
        c.font = header_font_m
        c.fill = header_fill_m
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_m

    for r_idx, rec in enumerate(records, start=4):
        vals = [
            rec['FY'], rec['Month'], rec['State'], rec['Vehicle Number'],
            rec['Make'], rec['Model'], rec['Year'], rec['Odometer'],
            rec['Final Bid Value'], rec['Seller Name'], rec['Seller Segment'],
            rec['Buyer Name'], rec['Total Income (Ex GST)'], rec['Total Income (Inc GST)'],
            f'=IF(AND(ISNUMBER(M{r_idx}), ISNUMBER(I{r_idx}), I{r_idx}>0), M{r_idx}/I{r_idx}, "")'
        ]
        for c_idx, val in enumerate(vals, start=1):
            cell = ws_master.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border_m
            if c_idx in (1, 2, 3):
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 7:
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "0"
            elif c_idx in (8,):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"
            elif c_idx in (9, 13, 14):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"
            elif c_idx == 15:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "0.00%"

    widths_m = [10, 10, 12, 18, 20, 20, 10, 15, 18, 36, 16, 34, 22, 18, 16]
    for i, w in enumerate(widths_m, start=1):
        ws_master.column_dimensions[get_column_letter(i)].width = w

    # Sheet 3: Standard Report
    ws_std = wb.create_sheet("Conversion Data (Standard)")
    ws_std.views.sheetView[0].showGridLines = True
    ws_std.freeze_panes = "A4"

    ws_std.merge_cells("A1:M1")
    ws_std["A1"] = "TVS Certified - Monthly Conversion Report (Standard)"
    ws_std["A1"].font = Font(name="Calibri", size=14, bold=True, color="2F5496")
    ws_std["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_std.merge_cells("A2:M2")
    ws_std["A2"] = f"{subtitle} - Standard 13-Column Manager Reporting Template"
    ws_std["A2"].font = Font(name="Calibri", size=10, italic=True, color="808080")
    ws_std["A2"].alignment = Alignment(horizontal="center", vertical="center")

    std_cols = [
        "FY", "Month", "State", "Vehicle Number", "Make", "Model",
        "Year", "Odometer", "Final Bid Value", "Seller Name",
        "Seller Segment", "Buyer Name", "Total Income"
    ]
    for c_i, h in enumerate(std_cols, start=1):
        c = ws_std.cell(row=3, column=c_i, value=h)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_m

    for r_idx, rec in enumerate(records, start=4):
        vals = [
            rec['FY'], rec['Month'], rec['State'], rec['Vehicle Number'],
            rec['Make'], rec['Model'], rec['Year'], rec['Odometer'],
            rec['Final Bid Value'], rec['Seller Name'], rec['Seller Segment'],
            rec['Buyer Name'], rec['Total Income (Ex GST)']
        ]
        for c_idx, val in enumerate(vals, start=1):
            cell = ws_std.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border_m
            if c_idx in (1, 2, 3):
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 7:
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "0"
            elif c_idx in (8, 9):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"
            elif c_idx == 13:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"

    widths_std = [10, 10, 12, 18, 20, 20, 10, 15, 18, 36, 16, 34, 20]
    for i, w in enumerate(widths_std, start=1):
        ws_std.column_dimensions[get_column_letter(i)].width = w

    # Sheet 4: _MakeList
    ws_makes = wb.create_sheet("_MakeList")
    ws_makes.sheet_state = "hidden"
    unique_makes = sorted(list(set(MAKE_MAP.values())))
    for i, m in enumerate(unique_makes, start=1):
        ws_makes.cell(row=i, column=1, value=m)

    safe_save(wb, output_path)


# =============================================================================
# WRITER 3: MASTER ENTRY TEMPLATE WITH FORMULAS & LOOKUP LISTS
# (Matches Conversion_Data_Template_Master.xlsx)
# =============================================================================

def write_master_entry_template(records, output_path):
    print(f"Writing Master Entry Template (16 Columns): {os.path.basename(output_path)}")
    wb = openpyxl.Workbook()

    # 1. Sheet 2: Lookup Lists
    ws_lookup = wb.create_sheet(title="Lookup Lists")
    ws_lookup.views.sheetView[0].showGridLines = True

    fy_list = ["FY24", "FY25", "FY26", "FY27", "FY28"]
    month_list = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    states_list = ["TN", "Kerala", "AP/TS", "Karnataka", "Maharashtra", "Other"]
    segment_list = ["Captive", "Non-Captive"]
    make_list = sorted(list(set(MAKE_MAP.values())))
    models_list = sorted(list(set(MODEL_MAP.values())))
    sellers_list = sorted(list(set([r['Seller Name'] for r in records if r.get('Seller Name')])))
    buyers_list = sorted(list(set([r['Buyer Name'] for r in records if r.get('Buyer Name')])))

    lookup_cols = [
        ("FY", fy_list),
        ("Month", month_list),
        ("State", states_list),
        ("Seller Segment", segment_list),
        ("Make", make_list),
        ("Model", models_list),
        ("Authorised Sellers", sellers_list),
        ("Verified Buyers", buyers_list),
    ]

    header_fill_l = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font_l = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    cell_font_l = Font(name="Calibri", size=10, color="0F172A")
    thin_border_l = Border(
        left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
    )

    lookup_ranges = {}
    for col_idx, (col_name, items) in enumerate(lookup_cols, start=1):
        col_letter = get_column_letter(col_idx)
        cell = ws_lookup.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill_l
        cell.font = header_font_l
        cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, val in enumerate(items, start=2):
            c = ws_lookup.cell(row=r_idx, column=col_idx, value=val)
            c.font = cell_font_l
            c.border = thin_border_l

        buffer_end = max(len(items) + 100, 500)
        lookup_ranges[col_name] = f"'Lookup Lists'!${col_letter}$2:${col_letter}${buffer_end}"
        ws_lookup.column_dimensions[col_letter].width = max(len(col_name) + 4, 18)

    # 2. Sheet 1: Main Data Entry Sheet
    ws_entry = wb.active
    ws_entry.title = "Conversion Data Entry"
    ws_entry.views.sheetView[0].showGridLines = True
    ws_entry.freeze_panes = "A2"

    columns_spec = [
        ("FY", 12, "0F172A"),
        ("Month", 12, "0F172A"),
        ("State", 14, "0F172A"),
        ("Vehicle Number", 18, "0F172A"),
        ("Make", 20, "0F172A"),
        ("Model", 22, "0F172A"),
        ("Seller Name", 34, "0F172A"),
        ("Seller Segment", 16, "0F172A"),
        ("Buyer Name", 32, "0F172A"),
        ("Year", 12, "0F172A"),
        ("Odometer", 16, "0F172A"),
        ("Final Bid Value", 18, "0F172A"),
        ("Total Income (Ex GST)", 22, "064E3B"),
        ("Total Income (Inc GST)", 24, "1E3A8A"),
        ("GST (18%)", 18, "312E81"),
        ("Period", 16, "334155"),
    ]

    for col_idx, (col_name, width, fill_hex) in enumerate(columns_spec, start=1):
        col_letter = get_column_letter(col_idx)
        cell = ws_entry.cell(row=1, column=col_idx, value=col_name)
        cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws_entry.column_dimensions[col_letter].width = width

    curr_format = "₹#,##0.00"
    num_format = "#,##0"
    int_format = "0"

    for r_idx, rec in enumerate(records, start=2):
        row_num = r_idx
        ws_entry.cell(row=r_idx, column=1, value=rec['FY']).alignment = Alignment(horizontal="center")
        ws_entry.cell(row=r_idx, column=2, value=rec['Month']).alignment = Alignment(horizontal="center")
        ws_entry.cell(row=r_idx, column=3, value=rec['State']).alignment = Alignment(horizontal="center")
        ws_entry.cell(row=r_idx, column=4, value=rec['Vehicle Number'])
        ws_entry.cell(row=r_idx, column=5, value=rec['Make'])
        ws_entry.cell(row=r_idx, column=6, value=rec['Model'])
        ws_entry.cell(row=r_idx, column=7, value=rec['Seller Name'])
        ws_entry.cell(row=r_idx, column=8, value=rec['Seller Segment'])
        ws_entry.cell(row=r_idx, column=9, value=rec['Buyer Name'])

        # Year
        c10 = ws_entry.cell(row=r_idx, column=10, value=rec['Year'])
        c10.number_format = int_format
        c10.alignment = Alignment(horizontal="center")

        # Odometer
        c11 = ws_entry.cell(row=r_idx, column=11, value=rec['Odometer'])
        c11.number_format = num_format
        c11.alignment = Alignment(horizontal="right")

        # Bid
        c12 = ws_entry.cell(row=r_idx, column=12, value=rec['Final Bid Value'])
        c12.number_format = curr_format
        c12.alignment = Alignment(horizontal="right")

        # Total Income Ex GST
        c13 = ws_entry.cell(row=r_idx, column=13, value=rec['Total Income (Ex GST)'])
        c13.number_format = curr_format
        c13.alignment = Alignment(horizontal="right")

        # Formulas
        c14 = ws_entry.cell(row=r_idx, column=14, value=f'=IF(M{row_num}="","",ROUND(M{row_num}*1.18,2))')
        c14.font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
        c14.number_format = curr_format
        c14.alignment = Alignment(horizontal="right")

        c15 = ws_entry.cell(row=r_idx, column=15, value=f'=IF(M{row_num}="","",ROUND(N{row_num}-M{row_num},2))')
        c15.font = Font(name="Calibri", size=10, color="312E81")
        c15.number_format = curr_format
        c15.alignment = Alignment(horizontal="right")

        c16 = ws_entry.cell(row=r_idx, column=16, value=f'=IF(AND(A{row_num}<>"",B{row_num}<>""), A{row_num}&" \u00b7 "&B{row_num}, "")')
        c16.font = Font(name="Calibri", size=10, color="475569")
        c16.alignment = Alignment(horizontal="center")

        for c_i in range(1, 17):
            c_cell = ws_entry.cell(row=r_idx, column=c_i)
            c_cell.border = thin_border_l
            if not c_cell.font.color:
                c_cell.font = cell_font_l

    # Dropdown data validations on Entry sheet
    max_row = max(len(records) + 200, 1000)

    for col_i, col_key in [(1, "FY"), (2, "Month"), (3, "State"), (5, "Make"), (6, "Model"), (7, "Authorised Sellers"), (8, "Seller Segment"), (9, "Verified Buyers")]:
        col_let = get_column_letter(col_i)
        rng = lookup_ranges[col_key]
        dv = DataValidation(type="list", formula1=f"={rng}", allow_blank=True)
        dv.add(f"{col_let}2:{col_let}{max_row}")
        ws_entry.add_data_validation(dv)

    safe_save(wb, output_path)


# =============================================================================
# UPDATE CUMULATIVE JAN-SEP DATASET
# =============================================================================

def update_cumulative_master(sept_records, jan_aug_path, output_paths, till_sept_path=None):
    print(f"Updating cumulative Jan-Sep master dataset from: {jan_aug_path}")
    df_hist = pd.read_excel(jan_aug_path, sheet_name=0, header=2)
    print(f"  Existing Jan-Aug rows: {len(df_hist)}")

    # Standardize historical columns
    df_hist_clean = df_hist[['FY', 'Month', 'State', 'Vehicle Number', 'Make', 'Model',
                             'Year', 'Odometer', 'Final Bid Value', 'Seller Name',
                             'Seller Segment', 'Buyer Name', 'Total Income (ex. GST)', 'Total Income']].copy()

    # Convert September records to matching DataFrame
    sept_rows = []
    for r in sept_records:
        sept_rows.append({
            'FY': r['FY'],
            'Month': r['Month'],
            'State': r['State'],
            'Vehicle Number': r['Vehicle Number'],
            'Make': r['Make'],
            'Model': r['Model'],
            'Year': r['Year'],
            'Odometer': r['Odometer'],
            'Final Bid Value': r['Final Bid Value'],
            'Seller Name': r['Seller Name'],
            'Seller Segment': r['Seller Segment'],
            'Buyer Name': r['Buyer Name'],
            'Total Income (ex. GST)': r['Total Income (Ex GST)'],
            'Total Income': r['Total Income (Inc GST)']
        })
    df_sept = pd.DataFrame(sept_rows)
    print(f"  September rows to append: {len(df_sept)}")

    # Combine
    df_combined = pd.concat([df_hist_clean, df_sept], ignore_index=True)
    print(f"  Total Combined Jan-Sep rows: {len(df_combined)}")

    # Prepare combined records dictionary list for write_analytics_workbook
    comb_records = []
    for _, row in df_combined.iterrows():
        ex = parse_num(row['Total Income (ex. GST)'], is_int=False)
        inc = parse_num(row['Total Income'], is_int=False)
        comb_records.append({
            'FY': str(row['FY']),
            'Month': str(row['Month']),
            'State': str(row['State']),
            'Vehicle Number': str(row['Vehicle Number']),
            'Make': str(row['Make']),
            'Model': str(row['Model']),
            'Year': parse_num(row['Year'], is_int=True),
            'Odometer': parse_num(row['Odometer'], is_int=True),
            'Final Bid Value': parse_num(row['Final Bid Value'], is_int=False),
            'Seller Name': str(row['Seller Name']) if pd.notna(row['Seller Name']) else "",
            'Seller Segment': str(row['Seller Segment']) if pd.notna(row['Seller Segment']) else "Non-Captive",
            'Buyer Name': str(row['Buyer Name']) if pd.notna(row['Buyer Name']) else "",
            'Total Income (Ex GST)': ex,
            'Total Income (Inc GST)': inc,
            'GST (18%)': round(inc - ex, 2) if (inc and ex) else None
        })

    for out_p in output_paths:
        write_analytics_workbook(comb_records, out_p, subtitle="January to September 2026")

    # If till_sept_path is provided, update sheet in 'Conversion Data till sept.xlsx'
    if till_sept_path and os.path.exists(till_sept_path):
        print(f"Updating 'Conversion Data till sept.xlsx'...")
        # Load existing workbook to preserve PIVOT sheet if present
        wb_till = openpyxl.load_workbook(till_sept_path)
        # We replace the data in 'Conversion Data FY26 - FY27'
        sheet_target = None
        for sname in wb_till.sheetnames:
            if 'conversion data' in sname.lower():
                sheet_target = sname
                break
        if sheet_target:
            ws_t = wb_till[sheet_target]
            # Clear existing data rows starting row 4
            for r in range(4, ws_t.max_row + 1):
                for c in range(1, 15):
                    ws_t.cell(row=r, column=c).value = None

            # Write combined records
            for r_idx, rec in enumerate(comb_records, start=4):
                vals = [
                    rec['FY'], rec['Month'], rec['State'], rec['Vehicle Number'],
                    rec['Make'], rec['Model'], rec['Year'], rec['Odometer'],
                    rec['Final Bid Value'], rec['Seller Name'], rec['Seller Segment'],
                    rec['Buyer Name'], rec['Total Income (Ex GST)']
                ]
                for c_idx, val in enumerate(vals, start=2): # Col B onwards in that sheet
                    cell = ws_t.cell(row=r_idx, column=c_idx, value=val)

            try:
                wb_till.save(till_sept_path)
                print(f"  -> Successfully updated: {till_sept_path}")
            except PermissionError:
                alt_path = os.path.splitext(till_sept_path)[0] + "_Updated.xlsx"
                wb_till.save(alt_path)
                print(f"  -> Original was open in Excel; saved updated version to: {alt_path}")


# =============================================================================
# MAIN FUNCTION
# =============================================================================

def main():
    print("=" * 75)
    print("  TVS CERTIFIED - SEPTEMBER 2026 CONVERSION DATA CONSOLIDATION")
    print("=" * 75)

    raw_sept_file = r"C:\Users\TC100\OneDrive - TVS Mobility Private Limited\Desktop\Conversion Data_Consolidation Automation\SEPT 2026.xlsx"
    dir_auto = r"C:\Users\TC100\OneDrive - TVS Mobility Private Limited\Desktop\Conversion Data_Consolidation Automation"
    dir_bi = r"c:\Users\TC100\OneDrive - TVS Mobility Private Limited\Desktop\BI Dashboard V2"

    if not os.path.exists(raw_sept_file):
        print(f"ERROR: File not found: {raw_sept_file}")
        sys.exit(1)

    # 1. Load and standardize
    sept_records = load_sept_data(raw_sept_file)

    # Summary
    st_counts = {}
    total_gmv = 0
    total_inc_ex = 0
    total_inc_inc = 0
    for r in sept_records:
        st = r['State']
        st_counts[st] = st_counts.get(st, 0) + 1
        total_gmv += (r['Final Bid Value'] or 0)
        total_inc_ex += (r['Total Income (Ex GST)'] or 0)
        total_inc_inc += (r['Total Income (Inc GST)'] or 0)

    print("\n--- Consolidated September Summary ---")
    for st, cnt in sorted(st_counts.items()):
        print(f"  {st:10}: {cnt} units")
    print(f"  TOTAL     : {len(sept_records)} units")
    print(f"  Turnover  : ₹{total_gmv:,.2f}")
    print(f"  Inc Ex-GST: ₹{total_inc_ex:,.2f}")
    print(f"  Inc Inc-GST: ₹{total_inc_inc:,.2f}")
    print(f"  GST (18%) : ₹{total_inc_inc - total_inc_ex:,.2f}\n")

    # 2. Output 1: Standard 13-column template format (September 2026)
    out_13col_auto = os.path.join(dir_auto, "Consolidated_Conversion_Data_Sep2026_Final.xlsx")
    write_standard_13col_template(sept_records, out_13col_auto)
    
    out_13col_bi = os.path.join(dir_bi, "Consolidated_Conversion_Data_Sep2026_Final.xlsx")
    write_standard_13col_template(sept_records, out_13col_bi)

    # 3. Output 2: Full Analytics Workbook with Executive Dashboard
    out_analytics_auto = os.path.join(dir_auto, "Consolidated_Conversion_Data_Sep2026_Analytics.xlsx")
    write_analytics_workbook(sept_records, out_analytics_auto, subtitle="September 2026")

    # 4. Output 3: Bulletproof Master Entry Template (16 Columns with formulas & dropdowns)
    out_entry_auto = os.path.join(dir_auto, "SEPT 2026_Consolidated_Master_Template.xlsx")
    write_master_entry_template(sept_records, out_entry_auto)

    out_entry_bi = os.path.join(dir_bi, "SEPT 2026_Consolidated_Master_Template.xlsx")
    write_master_entry_template(sept_records, out_entry_bi)

    # 5. Output 4: Update Cumulative Master Jan-Sep (3,791 + 523 = 4,314 units)
    jan_aug_master = os.path.join(dir_bi, "Conversion_Data_Jan_Aug2026_Analytics_GST Calculations.xlsx")
    cum_out_auto = os.path.join(dir_auto, "Consolidated_Conversion_Data_Jan_Sep2026_Analytics_GST Calculations.xlsx")
    cum_out_bi = os.path.join(dir_bi, "Conversion_Data_Jan_Sep2026_Analytics_GST Calculations.xlsx")
    till_sept_path = os.path.join(dir_auto, "Conversion Data till sept.xlsx")

    update_cumulative_master(
        sept_records=sept_records,
        jan_aug_path=jan_aug_master,
        output_paths=[cum_out_auto, cum_out_bi],
        till_sept_path=till_sept_path
    )

    print("\n" + "=" * 75)
    print("  ALL CONSOLIDATION TASKS COMPLETED SUCCESSFULLY!")
    print("=" * 75)

if __name__ == "__main__":
    main()
