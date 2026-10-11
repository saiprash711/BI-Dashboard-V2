import openpyxl
import os

f = r'c:\Users\TC100\OneDrive - TVS Mobility Private Limited\Desktop\BI Dashboard V2\SEPT 2026_Consolidated_Master_Template.xlsx'
wb = openpyxl.load_workbook(f)
ws_entry = wb['Conversion Data Entry']
ws_lookup = wb['Lookup Lists']

# Read lookup lists
lookups = {}
for c in range(1, ws_lookup.max_column + 1):
    h = ws_lookup.cell(1, c).value
    vals = set([ws_lookup.cell(r, c).value for r in range(2, ws_lookup.max_row + 1) if ws_lookup.cell(r, c).value is not None])
    lookups[h] = vals
    print(f'Lookup list "{h}": {len(vals)} items')

checks = [
    (1, 'FY', 'FY'),
    (2, 'Month', 'Month'),
    (3, 'State', 'State'),
    (5, 'Make', 'Make'),
    (6, 'Model', 'Model'),
    (7, 'Seller Name', 'Authorised Sellers'),
    (8, 'Seller Segment', 'Seller Segment'),
    (9, 'Buyer Name', 'Verified Buyers'),
]

total_missing = 0
for col_idx, col_name, lookup_key in checks:
    valid_set = lookups.get(lookup_key, set())
    missing = set()
    for r in range(2, ws_entry.max_row + 1):
        v = ws_entry.cell(r, col_idx).value
        if v and v not in valid_set:
            missing.add(v)
    print(f'Checking {col_name:15} against {lookup_key:20} -> Missing: {len(missing)}')
    if missing:
        print(f'   Unmatched samples: {list(missing)[:5]}')
        total_missing += len(missing)

print(f'\nTotal unmatched items across entire template: {total_missing}')
